# -*- coding: utf-8 -*-
import os
from pathlib import Path

docs = Path(r"D:\wamp64\www\kavminiapp\docs")
for f in docs.iterdir():
    print(repr(f.name), f.suffix, f.stat().st_size)

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "xlrd", "python-docx", "-q"])
    import openpyxl

import xlrd
from docx import Document

# Schools xlsx
schools = None
wards = None
consent = None
for f in docs.iterdir():
    n = f.name.lower()
    if f.suffix.lower() == ".xlsx" and "thanh" in n.replace("oà", "oa").replace("óa", "oa"):
        schools = f
    if "3321" in n and f.suffix.lower() in (".xls", ".xlsx"):
        wards = f
    if "consent" in n and f.suffix.lower() == ".docx":
        consent = f

# fallback search
if schools is None:
    for f in docs.glob("*.xlsx"):
        schools = f
        break
if wards is None:
    for f in docs.glob("*3321*"):
        wards = f
        break
if consent is None:
    for f in docs.glob("*.docx"):
        consent = f
        break

print("\n=== SCHOOLS FILE ===", schools)
wb = openpyxl.load_workbook(schools, read_only=True, data_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
print("sheets", wb.sheetnames)
print("header", rows[0] if rows else None)
print("sample rows:")
for r in rows[1:6]:
    print(r)
print("total data rows", max(0, len(rows) - 1))
# count non-empty
ids = []
for r in rows[1:]:
    if r and r[0] is not None:
        ids.append(r[0])
print("nonempty first-col", len(ids), "unique", len(set(ids)))

print("\n=== WARDS FILE ===", wards)
book = xlrd.open_workbook(str(wards))
print("sheets", book.sheet_names())
sh = book.sheet_by_index(0)
print("dims", sh.nrows, sh.ncols)
print("header", [sh.cell_value(0, c) for c in range(min(12, sh.ncols))])
for i in range(1, min(6, sh.nrows)):
    print([sh.cell_value(i, c) for c in range(min(12, sh.ncols))])

# find Thanh Hoa
th_count = 0
prov_col = None
for c in range(sh.ncols):
    h = str(sh.cell_value(0, c)).lower()
    if "tỉnh" in h or "tinh" in h or "province" in h:
        prov_col = c
        break
print("prov_col", prov_col)
if prov_col is not None:
    for i in range(1, sh.nrows):
        v = str(sh.cell_value(i, prov_col))
        if "Thanh" in v and "H" in v:
            th_count += 1
    print("Thanh Hoa-ish rows", th_count)

print("\n=== CONSENT DOCX ===", consent)
doc = Document(str(consent))
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        print(t)
for ti, table in enumerate(doc.tables):
    print(f"--- table {ti} ---")
    for row in table.rows:
        print(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
