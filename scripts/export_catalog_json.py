# -*- coding: utf-8 -*-
"""Export catalog Excel/XLS -> JSON for Laravel seeders."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
import xlrd

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT / "backend" / "database" / "data"
OUT.mkdir(parents=True, exist_ok=True)


def find_file(predicate):
    for f in DOCS.iterdir():
        if f.name.startswith(".~"):
            continue
        if predicate(f):
            return f
    raise FileNotFoundError(predicate)


def normalize_name(s: str) -> str:
    s = (s or "").replace("\n", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return s


def export_wards():
    path = find_file(lambda f: "3321" in f.name and f.suffix.lower() == ".xls")
    book = xlrd.open_workbook(str(path))
    sh = book.sheet_by_index(0)
    provinces = {}
    wards = []
    for i in range(1, sh.nrows):
        code = str(sh.cell_value(i, 0)).strip()
        if code.endswith(".0"):
            code = code[:-2]
        name = normalize_name(str(sh.cell_value(i, 1)))
        level = normalize_name(str(sh.cell_value(i, 2)))
        province_code = str(sh.cell_value(i, 4)).strip()
        if province_code.endswith(".0"):
            province_code = province_code[:-2]
        province_name = normalize_name(str(sh.cell_value(i, 5)))
        if not code or not name or not province_code:
            continue
        provinces[province_code] = province_name
        wards.append(
            {
                "code": code.zfill(5) if code.isdigit() else code,
                "name": name,
                "level": level,
                "province_code": province_code.zfill(2) if province_code.isdigit() else province_code,
                "province_name": province_name,
            }
        )
    (OUT / "wards_all.json").write_text(
        json.dumps({"provinces": provinces, "wards": wards}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    thanh_hoa = [w for w in wards if w["province_code"] in ("38", "038") or "Thanh H" in w["province_name"]]
    if not thanh_hoa:
        # fallback: code starting with 38
        thanh_hoa = [w for w in wards if str(w["province_code"]).lstrip("0") == "38" or w["province_code"] == "38"]
    th_code = thanh_hoa[0]["province_code"] if thanh_hoa else "38"
    th_name = thanh_hoa[0]["province_name"] if thanh_hoa else "Tỉnh Thanh Hóa"
    (OUT / "wards_thanh_hoa.json").write_text(
        json.dumps(
            {
                "province": {"code": th_code, "name": th_name},
                "wards": thanh_hoa,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    print(f"wards total={len(wards)} thanh_hoa={len(thanh_hoa)} province_code={th_code}")
    return th_code, {normalize_name(w["name"]).lower(): w for w in thanh_hoa}


def export_schools(ward_index):
    path = find_file(lambda f: f.suffix.lower() == ".xlsx" and "082026" in f.name)
    wb = load_workbook(path, read_only=True, data_only=True)
    # Prefer level sheets; fallback Toàn tỉnh
    preferred = []
    for name in wb.sheetnames:
        low = name.lower()
        if "tiểu" in low or "thcs" in low or "thpt" in low:
            preferred.append(name)
    sheets = preferred or [wb.sheetnames[0]]

    schools = []
    seen = set()
    unmatched_wards = {}

    for sheet_name in sheets:
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            continue
        # find columns
        headers = [str(h).strip().lower() if h is not None else "" for h in header]
        def col(*cands):
            for i, h in enumerate(headers):
                for c in cands:
                    if c in h:
                        return i
            return None

        id_col = col("id")
        ward_col = col("xã", "xa", "phường", "phuong")
        name_col = col("tên trường", "ten truong", "trường")
        # fallback fixed positions from inspect: STT, ID, Xã/phường, Tên trường
        if id_col is None:
            id_col = 1
        if ward_col is None:
            ward_col = 2
        if name_col is None:
            name_col = 3

        level = "unknown"
        low = sheet_name.lower()
        if "tiểu" in low:
            level = "tieu_hoc"
        elif "thcs" in low:
            level = "thcs"
        elif "thpt" in low or "gdtx" in low:
            level = "thpt_gdtx"

        for row in rows:
            if not row or len(row) <= max(id_col, ward_col, name_col):
                continue
            if row[id_col] is None:
                continue
            sid = str(row[id_col]).strip()
            if sid.endswith(".0"):
                sid = sid[:-2]
            if not sid.isdigit():
                continue
            ward_name = normalize_name(str(row[ward_col] or ""))
            school_name = normalize_name(str(row[name_col] or ""))
            if not school_name or not ward_name:
                continue
            key = sid
            if key in seen:
                continue
            seen.add(key)
            wkey = ward_name.lower()
            ward = ward_index.get(wkey)
            if not ward:
                # fuzzy: remove Phường/Xã/Thị trấn prefix
                alt = re.sub(r"^(phường|xã|thị trấn)\s+", "", wkey, flags=re.I)
                for kn, wv in ward_index.items():
                    kn2 = re.sub(r"^(phường|xã|thị trấn)\s+", "", kn, flags=re.I)
                    if alt == kn2 or alt in kn or kn in alt:
                        ward = wv
                        break
            if not ward:
                unmatched_wards[ward_name] = unmatched_wards.get(ward_name, 0) + 1
            schools.append(
                {
                    "external_id": sid,
                    "name": school_name,
                    "ward_name": ward_name,
                    "ward_code": ward["code"] if ward else None,
                    "level": level,
                    "sheet": sheet_name,
                }
            )

    (OUT / "schools_thanh_hoa.json").write_text(
        json.dumps(schools, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (OUT / "schools_unmatched_wards.json").write_text(
        json.dumps(unmatched_wards, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    matched = sum(1 for s in schools if s["ward_code"])
    print(f"schools={len(schools)} matched_ward={matched} unmatched_ward_names={len(unmatched_wards)}")
    if unmatched_wards:
        print("top unmatched:", sorted(unmatched_wards.items(), key=lambda x: -x[1])[:15])


def export_consent_text():
    from docx import Document

    path = find_file(lambda f: f.suffix.lower() == ".docx" and "consent" in f.name.lower())
    doc = Document(str(path))
    paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    (OUT / "consent_source_paragraphs.json").write_text(
        json.dumps(paras, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print("consent paragraphs", len(paras))


def main():
    th_code, ward_index = export_wards()
    export_schools(ward_index)
    export_consent_text()
    print("OUT", OUT)


if __name__ == "__main__":
    main()
